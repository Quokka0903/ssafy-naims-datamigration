# -*- coding: utf-8 -*-
"""담당자 검토용 Excel 워크북 생성."""
import io, os, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import build_mapping as B
import gen_docs as G

OUT = os.path.join(G.DOCS, "NAIMS_이관_매핑_검토본_%s.xlsx" % G.TAG)

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
BODY = Font(size=9)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

VERDICT_FILL = {
    "확정":      PatternFill("solid", fgColor="E2EFDA"),
    "자동생성":   PatternFill("solid", fgColor="D9E1F2"),
    "제외":      PatternFill("solid", fgColor="F2F2F2"),
    "미결":      PatternFill("solid", fgColor="FFF2CC"),
    "스키마변경": PatternFill("solid", fgColor="FCE4D6"),
    "신규제안":   PatternFill("solid", fgColor="DDEBF7"),
}


def sheet(wb, name, header, rows, widths, freeze="A2", fills=None):
    ws = wb.create_sheet(name)
    ws.append(header)
    for i, c in enumerate(ws[1], 1):
        c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY; c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=True)
    if fills:
        for ri, r in enumerate(rows, 2):
            f = fills(r)
            if f:
                for c in ws[ri]:
                    c.fill = f
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    return ws


def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── 표지 ──
    ws = wb.create_sheet("표지")
    ws["A1"] = "NAIMS → NAIMS2 데이터 이관 매핑 검토본"
    ws["A1"].font = TITLE_FONT
    meta = [
        ("문서번호", "NAIMS-MIG-%s" % G.TAG),
        ("작성일", G.STAMP),
        ("소스 DB (MsSQL)", G.SRC_MSSQL),
        ("소스 DB (Oracle)", G.SRC_ORACLE),
        ("대상 DB", G.TGT),
        ("기준 확정본", G.BASELINE),
        ("", ""),
        ("트랙① DB → DB", "대상 STD_* 전용 · %d개 테이블 (6단계) · 매핑 %d행" % (len(B.PLAN), len(G.M))),
        ("트랙② 파일 → DB", "대상 RAW_*/PREC_*/EXTRACT_DATA_STORE · 매핑 %d행" % len(G.FM)),
        ("", "RAW_*/PREC_*/EXT_* 는 DB→DB 이관 대상이 아님"),
        ("이관 제외 테이블", "%d개" % len(B.EXCLUDED)),
        ("매핑 총 행수", "%d행" % (len(G.M) + len(G.FM))),
        ("", ""),
        ("■ 검토 요청", "'결정요청' 시트의 [결정] 열에 판단을 기입해 주십시오."),
        ("", "전체 맥락은 '전컬럼매핑'(트랙①) / '파일이관매핑'(트랙②) 시트를 참조하시고,"),
        ("", "해당 시트의 [담당자검토] 열에도 자유롭게 의견을 남기실 수 있습니다."),
    ]
    r = 3
    for k, v in meta:
        ws.cell(r, 1, k).font = Font(bold=True, size=10)
        ws.cell(r, 2, v).font = Font(size=10)
        r += 1
    r += 1
    ws.cell(r, 1, "판정 구분").font = Font(bold=True, size=11)
    r += 1
    ws.cell(r, 2, "트랙① / 트랙② / 합계").font = Font(bold=True, size=9)
    r += 1
    for k, v in G.ALLVC.most_common():
        ws.cell(r, 1, k).fill = VERDICT_FILL.get(k)
        ws.cell(r, 1).font = Font(size=10)
        ws.cell(r, 2, "%d / %d / %d" % (G.VC.get(k, 0), G.FVC.get(k, 0), v)).font = Font(size=10)
        ws.cell(r, 3, {
            "확정": "기존 확정 + 현재 스키마 일치 — 그대로 구현",
            "제외": "이관하지 않음으로 확정",
            "미결": "대응 규칙 미확정 — 담당자 결정 필요",
            "스키마변경": "확정된 대상 컬럼이 현재 DB에 없음 — 재확정 필요",
            "신규제안": "규칙 기반 제안 — 승인 필요",
        }.get(k, "")).font = Font(size=10)
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 52

    # ── 이관순서 ──
    sheet(wb, "이관순서",
          ["단계", "소스 테이블", "대상 테이블", "소스 PK", "구분", "소스 레코드수"],
          sorted([[l, s, t, pk, c, G.ms_rows.get(s, "")] for s, t, pk, l, c in B.PLAN]),
          [8, 26, 30, 22, 12, 14])

    # ── 이관제외 ──
    sheet(wb, "이관제외", ["테이블", "구분", "사유", "MsSQL 레코드수"],
          [[k, "참조용" if k in B.REFERENCE_ONLY else "제외", v, G.ms_rows.get(k, "")]
           for k, v in sorted(B.EXCLUDED.items())],
          [26, 10, 46, 14])

    # ── 전컬럼매핑 ──
    rows = [[r["cat"], r["level"], r["src"], r["col"], r["mtyp"], r["otyp"],
             r["tgt"], r["tcol"], r["ttyp"], r["rule"], r["verdict"], r["note"], ""]
            for r in G.M]
    sheet(wb, "전컬럼매핑",
          ["구분", "단계", "소스 테이블", "소스 컬럼", "MsSQL 타입", "Oracle 타입",
           "대상 테이블", "대상 컬럼", "대상 타입", "변환 규칙", "판정", "비고", "담당자검토"],
          rows, [10, 6, 22, 24, 16, 18, 26, 26, 16, 22, 10, 44, 30],
          freeze="D2", fills=lambda r: VERDICT_FILL.get(r[10]))

    # ── 파일이관매핑 (트랙②) ──
    frows = [[r["kind"], r["ent"], r["pat"], r["obj"], r["fld"],
              r["tgt"], r["tcol"], r["ttyp"], r["rule"], r["verdict"], r["note"], ""]
             for r in G.FM]
    sheet(wb, "파일이관매핑",
          ["구분", "소스 파일", "경로 패턴", "파싱 객체", "소스 항목",
           "대상 테이블", "대상 컬럼", "대상 타입", "변환 규칙", "판정", "비고", "담당자검토"],
          frows, [12, 22, 34, 20, 30, 26, 28, 16, 34, 10, 40, 30],
          freeze="E2", fills=lambda r: VERDICT_FILL.get(r[9]))

    # ── 결정요청 (양 트랙 통합) ──
    dec = [["① DB→DB", r["verdict"], r["src"], r["col"], r["mtyp"],
            r["tgt"], r["tcol"], r["note"], "", "", ""]
           for r in G.M if r["verdict"] in ("미결", "스키마변경", "신규제안")]
    dec += [["② 파일→DB", r["verdict"], r["ent"], r["fld"], "-",
             r["tgt"], r["tcol"], r["note"] or r["rule"], "", "", ""]
            for r in G.FM if r["verdict"] in ("미결", "스키마변경", "신규제안")]
    sheet(wb, "결정요청",
          ["트랙", "판정", "소스 테이블/파일", "소스 컬럼/항목", "소스 타입",
           "대상 테이블", "대상 컬럼", "사유·비고", "결정(기입)", "담당자", "결정일"],
          dec, [12, 12, 24, 26, 16, 26, 26, 50, 26, 12, 12],
          fills=lambda r: VERDICT_FILL.get(r[1]))

    # ── 소스구조 비교 ──
    O, Ms = set(G.ora), set(G.ms) - {"sysdiagrams"}
    cmp_rows = []
    for tb in sorted(O | Ms):
        oc = [c for c, _ in G.ora.get(tb, [])]
        mc = [c for c, _ in G.ms.get(tb, [])]
        if tb not in O:   note = "MsSQL에만 존재"
        elif tb not in Ms: note = "Oracle에만 존재"
        else:
            d = [c for c in mc if c not in oc]
            note = ("컬럼 차이: " + ", ".join(d)) if d else "동일"
        cmp_rows.append([tb, len(oc) or "-", len(mc) or "-", note,
                         "이관대상" if tb in G.PLAN_TABLES else ("제외" if tb in B.EXCLUDED else "-")])
    sheet(wb, "소스구조비교",
          ["테이블", "Oracle 컬럼수", "MsSQL 컬럼수", "차이", "이관 여부"],
          cmp_rows, [28, 14, 14, 46, 12])

    # ── 소스구조 상세 ──
    det = []
    for tb in sorted(set(G.ms) | set(G.ora)):
        od, md = dict(G.ora.get(tb, [])), dict(G.ms.get(tb, []))
        for c in [x for x, _ in G.ms.get(tb, [])] + \
                 [x for x, _ in G.ora.get(tb, []) if x not in md]:
            det.append([tb, c, md.get(c, "-"), od.get(c, "-")])
    sheet(wb, "소스구조상세", ["테이블", "컬럼", "MsSQL 타입", "Oracle 타입"],
          det, [28, 28, 20, 22])

    # ── 대상구조 (이관 대상 테이블만) ──
    tgt_rows = []
    for s, t, pk, l, c in B.PLAN:
        for col, ty in G.tg.get(t, []):
            tgt_rows.append([t, col, ty, s])
    sheet(wb, "대상구조", ["대상 테이블", "컬럼", "타입", "대응 소스 테이블"],
          tgt_rows, [30, 30, 22, 24])

    # ── 결정규칙 ──
    sheet(wb, "기존결정규칙", ["구분", "내용"],
          [[d.get("구분", ""), d.get("내용", "")] for d in G.DEC], [14, 100])

    wb.save(OUT)
    print("saved: %s" % OUT)
    print("sheets: %s" % ", ".join(wb.sheetnames))
    print("mapping rows: %d / 결정요청: %d" % (len(G.M), len(dec)))


if __name__ == "__main__":
    main()
